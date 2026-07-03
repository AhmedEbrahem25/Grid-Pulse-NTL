"""Generate Grid-Pulse-NTL-Formal.pptx — a formal, light/navy investor deck
(prospectus register), matching pitch-formal/index.html. Embeds the real kiosk +
dashboard photos. Run in a container with python-pptx:  python build_pptx.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PAPER=RGBColor(0xF7,0xF8,0xF9); WHITE=RGBColor(0xFF,0xFF,0xFF)
INK=RGBColor(0x10,0x23,0x3A); NAVY=RGBColor(0x0F,0x2A,0x4A); ACCENT=RGBColor(0x1E,0x6F,0x7A)
MUTED=RGBColor(0x5B,0x6B,0x7A); LINE=RGBColor(0xD9,0xDE,0xE3); CRIT=RGBColor(0xA3,0x24,0x2B)
SERIF="Georgia"; SANS="Segoe UI"; MONO="Consolas"
IMG=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","assets")

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def slide(bg=PAPER):
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    return s

def box(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    return tf

def para(tf,runs,size=14,color=INK,bold=False,font=SANS,align=PP_ALIGN.LEFT,space_after=4,space_before=0,line=1.2,first=False,italic=False):
    p=tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment=align;p.space_after=Pt(space_after);p.space_before=Pt(space_before)
    try:p.line_spacing=line
    except Exception:pass
    if isinstance(runs,str):runs=[(runs,{})]
    for text,ov in runs:
        r=p.add_run();r.text=text;f=r.font
        f.size=Pt(ov.get("size",size));f.bold=ov.get("bold",bold);f.name=ov.get("font",font)
        f.italic=ov.get("italic",italic);f.color.rgb=ov.get("color",color)
        if ov.get("spacing"):
            r._r.get_or_add_rPr().set("spc",str(int(ov["spacing"]*100)))
    return p

def rect(s,x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.fill.background();sh.shadow.inherit=False
    return sh

def card(s,x,y,w,h,fill=WHITE,line=LINE):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.color.rgb=line;sh.line.width=Pt(0.75);sh.shadow.inherit=False
    return sh

def photo(s,name,x,y,w):
    p=s.shapes.add_picture(os.path.join(IMG,name),Inches(x),Inches(y),width=Inches(w))
    p.line.color.rgb=LINE;p.line.width=Pt(1);p.shadow.inherit=False
    return p

def eyebrow(s,x,y,num,label):
    tf=box(s,x,y,11,0.4)
    para(tf,[(f"{num}    ",{"color":ACCENT,"font":MONO,"size":13}),
             (label,{"color":ACCENT,"font":SANS,"size":11.5,"spacing":2.2,"bold":True})],first=True)

def title(s,t):
    rect(s,0.7,1.02,0.5,0.03,ACCENT)
    tf=box(s,0.7,1.12,11.9,1.0)
    para(tf,[(t,{"color":NAVY,"size":30,"bold":True,"font":SERIF})],first=True)

def lead(s,t,y=2.05,w=11.6):
    tf=box(s,0.7,y,w,0.8)
    para(tf,[(t,{"color":INK,"size":15})],line=1.35,first=True)

# ===================================================== 1 cover
s=slide()
rect(s,0,0,13.333,0.14,NAVY)
tf=box(s,0.8,0.55,12,0.4)
para(tf,[("CONFIDENTIAL",{"color":CRIT,"font":SANS,"size":11.5,"spacing":1.8,"bold":True}),
         ("   ·   INVESTMENT & DEPLOYMENT PROPOSAL · 2026",{"color":MUTED,"font":SANS,"size":11.5,"spacing":1.8})],first=True)
tf=box(s,0.8,1.7,11.6,2.2)
para(tf,[("Autonomous Detection of Non-Technical Loss",{"color":NAVY,"size":36,"bold":True,"font":SERIF})],line=1.08,first=True)
para(tf,[("on the Egyptian Distribution Grid",{"color":NAVY,"size":36,"bold":True,"font":SERIF})],line=1.08)
tf=box(s,0.8,3.9,10.6,0.9)
para(tf,[("An AI edge-and-cloud system that distinguishes electricity theft from the natural losses of heat and load — and localizes it to the transformer, with explainable evidence.",
          {"color":MUTED,"size":16,"italic":True,"font":SERIF})],line=1.35,first=True)
kf=[("EGP 25–35B","annual non-technical loss (Egypt)",CRIT),("94%","theft-detection confidence, demonstrated",NAVY),
    ("< 2 months","utility payback per node",NAVY),("EGP 30M","seed investment sought",NAVY)]
for i,(v,d,col) in enumerate(kf):
    x=0.8+i*3.0; card(s,x,5.35,2.85,1.35)
    tf=box(s,x+0.22,5.55,2.5,0.5);para(tf,[(v,{"color":col,"font":MONO,"size":18,"bold":True})],first=True)
    tf=box(s,x+0.22,6.05,2.5,0.6);para(tf,[(d,{"color":MUTED,"size":10.5})],line=1.25,first=True)

def section(numstr, ttl):
    s=slide(); eyebrow(s,0.7,0.6,numstr,"GRID-PULSE NTL · FORMAL PROPOSAL"); title(s,ttl); return s

# ===================================================== 2 executive summary
s=section("01","Executive Summary")
tf=box(s,0.7,2.0,11.6,4.5)
para(tf,[("Grid-Pulse NTL is a business-to-government venture that recovers stolen electricity for distribution companies. A low-cost sensor on each distribution transformer measures the energy leaving it; the platform compares that against registered consumption and an AI-predicted technical (natural) loss, and reports the unexplained excess as suspected theft — with a probability, a reasons checklist, and an estimate of the kilowatts at risk.",{"color":INK,"size":15})],line=1.4,first=True)
para(tf,[("The commercial model aligns the company with the utility: a share of the energy recovered, atop a modest hardware and platform fee. At typical recovery rates a single node pays for itself in under two months. The system is built and demonstrated.",{"color":INK,"size":15})],line=1.4,space_before=10)
card(s,0.7,5.35,11.9,1.5,fill=RGBColor(0xEE,0xF3,0xF4),line=RGBColor(0xEE,0xF3,0xF4))
rect(s,0.7,5.35,0.05,1.5,ACCENT)
tf=box(s,1.0,5.55,11.3,1.2)
para(tf,[("The differentiator is not measurement — it is attribution. ",{"color":NAVY,"size":14.5,"bold":True}),
         ("An AI baseline predicts how much loss is normal for the present heat and load, so only the genuinely unexplained remainder is flagged. This removes the false alarms that make simple meter-difference methods unusable in the Egyptian summer.",{"color":INK,"size":14.5})],line=1.35,first=True)

# ===================================================== 3 problem
s=section("02","The Problem: Non-Technical Loss")
tf=box(s,0.7,2.0,11.6,4)
para(tf,[("Distribution companies lose energy two ways. ",{"color":INK,"size":15}),
         ("Technical loss",{"color":NAVY,"size":15,"bold":True}),
         (" is physical — heat in cables and transformers — and rises with temperature and load. ",{"color":INK,"size":15}),
         ("Non-technical loss (NTL)",{"color":NAVY,"size":15,"bold":True}),
         (" is theft: illegal hooks, tampered meters, unmetered workshops and kilns. In Egypt, NTL is estimated at EGP 25–35 billion per year, and it worsens load-shedding by overloading transformers with demand no one is billed for.",{"color":INK,"size":15})],line=1.4,first=True)
para(tf,[("The operational difficulty: both losses present identically at the transformer, as “more loss.” A fixed threshold alarms on every hot afternoon and is ignored — leaving utilities with slow, tip-driven manual inspection.",{"color":INK,"size":15})],line=1.4,space_before=12)

# ===================================================== 4 solution (+ dashboard)
s=section("03","The Solution")
tf=box(s,0.7,2.0,5.2,4.5)
para(tf,[("Grid-Pulse runs a continuous energy balance at each transformer, predicts the expected technical loss for current conditions, and treats the difference as the candidate theft signal.",{"color":INK,"size":14.5})],line=1.4,first=True)
para(tf,[("That signal is confirmed or dampened by three independent lines of evidence — harmonic signature, power factor, and a longer-horizon energy accounting — before a verdict is issued.",{"color":INK,"size":14.5})],line=1.4,space_before=10)
para(tf,[("Every verdict is explainable: a probability, a checklist of the indicators that fired, the kilowatts unaccounted for, and a recommended action.",{"color":INK,"size":14.5})],line=1.4,space_before=10)
photo(s,"dashboard.jpg",6.4,1.95,6.2)
tf=box(s,6.4,7.0,6.2,0.4)
para(tf,[("Figure 1. ",{"color":NAVY,"size":10.5,"bold":True}),("The operator console — a demonstrated 94% theft verdict.",{"color":MUTED,"size":10.5})],first=True)

# ===================================================== 5 technology + eval table
s=section("04","Technology & Detection")
layers=[("Energy balance","measured output less registered consumption"),
        ("AI technical-loss baseline","Random-Forest over 7 features → expected natural loss"),
        ("Harmonic signature (1D-CNN)","clean / legit-industrial / illegal-bypass"),
        ("Transformer energy accounting","excess loss over a horizon as corroborating evidence"),
        ("Temporal anomaly","Isolation-Forest on the consumption pattern"),
        ("Decision fusion & XAI","calibrated probability + reasons checklist")]
yy=2.0
for t,d in layers:
    tf=box(s,0.7,yy,6.0,0.55)
    para(tf,[("•  ",{"color":ACCENT,"size":13,"bold":True}),(t,{"color":NAVY,"size":13.5,"bold":True})],first=True)
    para(tf,[("     "+d,{"color":MUTED,"size":11.5})],space_before=1)
    yy+=0.72
# eval table right
tf=box(s,7.3,1.95,5.2,0.4);para(tf,[("MEASURED PERFORMANCE (held-out, synthetic)",{"color":ACCENT,"font":SANS,"size":10.5,"spacing":1.2,"bold":True})],first=True)
rows=[("Theft precision","1.00"),("Theft recall","0.78"),("F1 score","0.88"),("Brier score","0.06"),("Legit-industrial false alarm","0.00")]
yy=2.5
for k,v in rows:
    strong = k.startswith("Legit")
    tf=box(s,7.3,yy,3.9,0.32,anchor=MSO_ANCHOR.MIDDLE);para(tf,[(k,{"color":(NAVY if strong else INK),"size":12.5,"bold":strong})],first=True)
    tf=box(s,11.3,yy,1.2,0.32,anchor=MSO_ANCHOR.MIDDLE);para(tf,[(v,{"color":(ACCENT if strong else NAVY),"font":MONO,"size":13,"bold":True})],align=PP_ALIGN.RIGHT,first=True)
    rect(s,7.3,yy+0.4,5.2,0.012,LINE);yy+=0.52
tf=box(s,7.3,yy+0.1,5.2,0.5);para(tf,[("Models retrain on real utility meter data in production.",{"color":MUTED,"size":10.5,"italic":True})],first=True)

# ===================================================== 6 deployment (+ kiosk)
s=section("05","Field Deployment")
photo(s,"kiosk.jpg",0.7,2.0,3.7)
tf=box(s,0.7,5.3,3.7,0.4);para(tf,[("Figure 2. ",{"color":NAVY,"size":10.5,"bold":True}),("Distribution kiosk (RMU), Kafr El-Sheikh.",{"color":MUTED,"size":10.5})],first=True)
tf=box(s,5.0,2.0,7.6,4.5)
para(tf,[("The field target is a sealed distribution kiosk (ring-main unit). Grid-Pulse retrofits it non-invasively, on the low-voltage side only:",{"color":INK,"size":15})],line=1.4,first=True)
for b in ["Three Rogowski coils clipped around the LV busbars — one per phase.",
          "An edge node in a weatherproof IP65 enclosure; antenna routed outside the shell.",
          "The 11 kV chamber is never entered; no conductor cut; no feeder disconnected.",
          "Fully reversible; ~30 minutes per kiosk; no service interruption.",
          "Production: Rogowski coils + LoRaWAN/cellular, suitable for ~230,000 transformers."]:
    para(tf,[("—  ",{"color":ACCENT,"size":13}),(b,{"color":INK,"size":14})],space_before=8,line=1.35)

# ===================================================== 7 market
s=section("06","Market & Competitive Position")
tf=box(s,0.7,2.0,11.8,4)
para(tf,[("Utilities fight theft today by manual audit, full smart-meter rollout, analytics on meter data, feeder monitoring, or fixed-threshold alarms. Each is slow, blind to cause, or dependent on a metering estate that costs billions to build.",{"color":INK,"size":15})],line=1.4,first=True)
para(tf,[("Grid-Pulse occupies the position none of these hold: low capital cost, explainable, and effective on a grid that was never fully metered. ",{"color":NAVY,"size":15,"bold":True}),
         ("One node observes a whole feeder of customers — a fraction of per-meter cost — and the AI baseline removes the false positives that make cheaper analytics unusable. Each verdict adds labelled theft data, a compounding advantage a fixed-threshold competitor cannot replicate.",{"color":INK,"size":15})],line=1.4,space_before=12)

# ===================================================== 8 financials
s=section("07","Business Model & Financials")
tf=box(s,0.7,1.95,6,0.35);para(tf,[("REVENUE STREAMS",{"color":ACCENT,"font":SANS,"size":10.5,"spacing":1.4,"bold":True})],first=True)
for t,d in [("Hardware — EGP 6,000 / node","edge node + gateway share, at margin or leased"),
            ("Platform — EGP 300 / node / month","detection, dashboard, verdicts, reporting"),
            ("Recovery share — 15%","of verified recovered energy, first 24 months")]:
    yy=2.35+[("Hardware — EGP 6,000 / node"),("Platform — EGP 300 / node / month"),("Recovery share — 15%")].index(t)*0.72
    tf=box(s,0.7,yy,6,0.6)
    para(tf,[(t,{"color":NAVY,"size":13,"bold":True})],first=True)
    para(tf,[(d,{"color":MUTED,"size":11.5})],space_before=1)
tf=box(s,0.7,4.7,6,1.0);para(tf,[("Unit economics: ",{"color":NAVY,"size":13,"bold":True}),("5-yr gross profit ≈ EGP 20,750 / node (LTV/CAC ≈ 8×); customer payback < 2 months.",{"color":INK,"size":13})],line=1.35,first=True)
# 5-yr table right
cols=["EGP M","Y1","Y2","Y3","Y4","Y5"]
tf=box(s,7.0,1.95,5.6,0.3)
p=tf.paragraphs[0]
for i,c in enumerate(cols):
    r=p.add_run();r.text=(c if i==0 else c.rjust(7));r.font.name=MONO;r.font.size=Pt(11);r.font.color.rgb=ACCENT;r.font.bold=True
data=[("Total revenue",["4.4","33.5","142.8","324.4","604.8"],True),
      ("Gross profit",["1.9","15.0","67.0","167.1","327.8"],False),
      ("Op. expenses",["12.0","28.0","50.0","95.0","160.0"],False),
      ("EBITDA",["(10.1)","(13.0)","17.0","72.1","167.8"],True)]
yy=2.4
for name,vals,strong in data:
    tf=box(s,7.0,yy,2.3,0.32,anchor=MSO_ANCHOR.MIDDLE);para(tf,[(name,{"color":NAVY if strong else INK,"size":12,"bold":strong})],first=True)
    for i,v in enumerate(vals):
        tf=box(s,9.2+i*0.72,yy,0.68,0.32,anchor=MSO_ANCHOR.MIDDLE)
        col=CRIT if v.startswith("(") else NAVY
        para(tf,[(v,{"color":col,"font":MONO,"size":11.5,"bold":strong})],align=PP_ALIGN.RIGHT,first=True)
    rect(s,7.0,yy+0.4,5.6,0.012,LINE);yy+=0.52
tf=box(s,7.0,yy+0.15,5.6,0.5);para(tf,[("Founder's model; breakeven in Year 3.",{"color":MUTED,"size":10.5,"italic":True})],first=True)

# ===================================================== 9 ask + risk
s=section("08","The Investment Ask")
tf=box(s,0.7,2.0,5.4,2.4)
para(tf,[("EGP 30M",{"color":NAVY,"font":SERIF,"size":42,"bold":True})],first=True)
para(tf,[("≈ USD 625,000 seed — through a district pilot and two years to the Year-3 EBITDA turn.",{"color":INK,"size":14})],line=1.35,space_before=8)
para(tf,[("Use of funds: field pilot 40% · product & ML 25% · hardware R&D + certification 15% · team & G&A 12% · go-to-market 8%.",{"color":MUTED,"size":12.5})],line=1.35,space_before=10)
tf=box(s,6.4,2.0,6.2,0.35);para(tf,[("RISK & GOVERNANCE",{"color":ACCENT,"font":SANS,"size":10.5,"spacing":1.4,"bold":True})],first=True)
for t,m in [("Procurement cycles","paid pilot + audited recovery → framework tender"),
            ("Hardware capital","sold at margin; fleets asset-financed vs contracts"),
            ("Model trust","explainable, human-reviewed, retrained on real data"),
            ("Data & privacy","meters transformers, not households")]:
    yy=2.45+[("Procurement cycles"),("Hardware capital"),("Model trust"),("Data & privacy")].index(t)*0.72
    tf=box(s,6.4,yy,6.2,0.6)
    para(tf,[("—  ",{"color":ACCENT,"size":12}),(t,{"color":NAVY,"size":13,"bold":True})],first=True)
    para(tf,[("     "+m,{"color":MUTED,"size":11.5})],space_before=1)
rect(s,0.7,7.0,12.0,0.02,NAVY)
tf=box(s,0.7,7.05,12,0.3);para(tf,[("Confidential · Grid-Pulse NTL · Investment & Deployment Proposal · 2026",{"color":MUTED,"size":9.5})],first=True)

out="Grid-Pulse-NTL-Formal.pptx";prs.save(out)
print("saved",out,"-",len(prs.slides._sldIdLst),"slides")
