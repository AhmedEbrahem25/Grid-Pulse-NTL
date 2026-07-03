"""Generate Grid-Pulse-NTL-Business.pptx — a 16:9 investor deck for the business
& financial model, matching the pitch identity (petrol-ink, signal-cyan, mono
readouts) and embedding the real kiosk photo. Run in a container with python-pptx:

    pip install python-pptx && python build_pptx.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK=RGBColor(0x0A,0x0E,0x0F); INK2=RGBColor(0x10,0x16,0x1A); PANEL=RGBColor(0x0E,0x14,0x18)
LINE=RGBColor(0x1E,0x2A,0x2F); LINE2=RGBColor(0x2A,0x3A,0x41); FG=RGBColor(0xE8,0xEE,0xF0)
MUTED=RGBColor(0x8A,0x9B,0xA1); DIM=RGBColor(0x5F,0x71,0x78); SIGNAL=RGBColor(0x34,0xC8,0xE0)
SIGDIM=RGBColor(0x1C,0x6A,0x74); THEFT=RGBColor(0xFF,0x4D,0x5E); TECH=RGBColor(0xF4,0xA6,0x3B)
OK=RGBColor(0x3D,0xDC,0x84); CY=RGBColor(0xC3,0xD0,0xD4)
SANS="Segoe UI"; SANS_B="Segoe UI Semibold"; MONO="Consolas"
IMG=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","assets")

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def slide():
    s=prs.slides.add_slide(BLANK)
    bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb=INK; bg.line.fill.background(); bg.shadow.inherit=False
    return s

def box(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    return tf

def para(tf,runs,size=14,color=FG,bold=False,font=SANS,align=PP_ALIGN.LEFT,space_after=4,space_before=0,line=1.12,first=False):
    p=tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space_after); p.space_before=Pt(space_before)
    try: p.line_spacing=line
    except Exception: pass
    if isinstance(runs,str): runs=[(runs,{})]
    for text,ov in runs:
        r=p.add_run(); r.text=text; f=r.font
        f.size=Pt(ov.get("size",size)); f.bold=ov.get("bold",bold); f.name=ov.get("font",font)
        f.color.rgb=ov.get("color",color)
        if ov.get("spacing"):
            rPr=r._r.get_or_add_rPr(); rPr.set("spc",str(int(ov["spacing"]*100)))
    return p

def rrect(s,x,y,w,h,fill=PANEL,line=LINE,line_w=1.0,radius=0.06):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(line_w)
    sh.shadow.inherit=False
    try: sh.adjustments[0]=radius
    except Exception: pass
    return sh

def rect(s,x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False
    return sh

def strip(s,x,y,h,color):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(0.045),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=color; r.line.fill.background(); r.shadow.inherit=False

def eyebrow(s,x,y,num,label):
    tf=box(s,x,y,11,0.4)
    para(tf,[(f"{num}   ",{"color":SIGDIM,"font":MONO,"size":13}),
             (label,{"color":SIGNAL,"font":MONO,"size":12.5,"spacing":2.4})],first=True)

def title(s,t):
    tf=box(s,0.7,1.02,11.9,1.0)
    para(tf,[(t,{"color":FG,"size":33,"bold":True,"font":SANS_B})],first=True)

def kicker(s,t):
    tf=box(s,0.7,1.9,11.6,0.7)
    para(tf,[(t,{"color":MUTED,"size":14.5})],line=1.3,first=True)

def photo(s,name,x,y,w):
    pic=s.shapes.add_picture(os.path.join(IMG,name),Inches(x),Inches(y),width=Inches(w))
    pic.line.color.rgb=LINE2; pic.line.width=Pt(1.25); pic.shadow.inherit=False
    return pic

# ===================================================== 1 · title
s=slide()
tf=box(s,0.7,0.55,12,0.4)
para(tf,[("GRID-PULSE NTL",{"color":SIGNAL,"font":MONO,"size":13,"spacing":3.0,"bold":True}),
         ("   //  Business & Financial Model",{"color":DIM,"font":MONO,"size":12})],first=True)
tf=box(s,0.7,1.6,10.6,2.0)
para(tf,[("The unit economics of ",{"color":FG,"size":38,"bold":True,"font":SANS_B}),
         ("catching theft",{"color":SIGNAL,"size":38,"bold":True,"font":SANS_B}),
         (" pay for themselves in weeks.",{"color":FG,"size":38,"bold":True,"font":SANS_B})],line=1.05,first=True)
tf=box(s,0.7,3.9,10.8,1.0)
para(tf,[("A B2G venture that recovers stolen electricity for Egypt's distribution companies — hardware on the "
          "transformer, AI in the cloud, and a share of every kilowatt it puts back on the bill.",{"color":CY,"size":15})],line=1.35,first=True)
cards=[("EGP 48 = $1","reference rate (2026)"),("~230,000","distribution transformers (est.)"),
       ("EGP 25–35B","annual non-technical loss"),("< 2 months","utility payback per node")]
for i,(v,d) in enumerate(cards):
    x=0.7+i*3.05; rrect(s,x,5.35,2.85,1.35,fill=PANEL,line=LINE,radius=0.08)
    tf=box(s,x+0.22,5.55,2.5,0.5); para(tf,[(v,{"color":FG,"font":MONO,"size":19,"bold":True})],first=True)
    tf=box(s,x+0.22,6.05,2.5,0.6); para(tf,[(d,{"color":MUTED,"size":11})],line=1.2,first=True)

# ===================================================== 2 · business model canvas
s=slide(); eyebrow(s,0.7,0.6,"01","BUSINESS MODEL"); title(s,"How the money flows")
blocks=[("CUSTOMER SEGMENTS",SIGNAL,["Distribution companies (EDCs) — primary","Ministry of Electricity / regulator","Industrial estates; MENA utilities next"]),
        ("VALUE PROPOSITION",SIGNAL,["Recover stolen energy — new revenue, not tariffs","Separate theft from natural loss → few false alarms","Explainable verdict an inspector acts on"]),
        ("REVENUE STREAMS",OK,["Hardware sale / HaaS lease","SaaS platform subscription","Performance / recovery share"]),
        ("KEY RESOURCES · MOAT",TECH,["Labelled theft dataset (compounds)","Edge hardware + detection IP","Ministry & field-ops relationships"]),
        ("CHANNELS",SIGNAL,["Direct BD to EDCs & Ministry","Paid pilots → framework tenders","System integrators"]),
        ("COST STRUCTURE",MUTED,["Hardware COGS & logistics","Cloud + ML compute","Field deployment & connectivity"])]
for i,(h,col,items) in enumerate(blocks):
    x=0.7+(i%3)*4.03; y=2.55+(i//3)*2.15
    rrect(s,x,y,3.85,1.95,fill=PANEL,line=LINE,radius=0.05)
    tf=box(s,x+0.24,y+0.2,3.4,0.3); para(tf,[(h,{"color":col,"font":MONO,"size":10,"spacing":1.0})],first=True)
    tf=box(s,x+0.24,y+0.55,3.45,1.3)
    for j,it in enumerate(items):
        para(tf,[("– ",{"color":SIGDIM}),(it,{"color":CY,"size":11.5})],first=(j==0),space_after=4,line=1.18)

# ===================================================== 3 · revenue streams
s=slide(); eyebrow(s,0.7,0.6,"02","REVENUE"); title(s,"Three revenue streams")
kicker(s,"Lead with a hardware sale + a light SaaS fee, then capture upside on the energy we actually recover.")
streams=[("STREAM 01 · one-time","Hardware","EGP 6,000","/ node",SIGNAL,"Edge node + gateway share at ~25% margin. Or HaaS lease to shift CapEx off the utility."),
         ("STREAM 02 · recurring","SaaS platform","EGP 300","/ node / month",OK,"Detection, live map, explainable verdicts, reporting. ~75% gross margin."),
         ("STREAM 03 · performance","Recovery share","15%","of recovered energy",TECH,"A cut of verified recovered billing for 24 months per zone. We win only when they do.")]
for i,(tg,h,price,unit,col,desc) in enumerate(streams):
    x=0.7+i*4.03; rrect(s,x,2.7,3.85,3.6,fill=PANEL,line=LINE,radius=0.05); strip(s,x,2.7,3.6,col)
    tf=box(s,x+0.28,2.95,3.4,0.3); para(tf,[(tg,{"color":col,"font":MONO,"size":10.5,"spacing":0.6})],first=True)
    tf=box(s,x+0.28,3.35,3.4,0.4); para(tf,[(h,{"color":FG,"size":19,"bold":True})],first=True)
    tf=box(s,x+0.28,3.95,3.4,0.6); para(tf,[(price,{"color":FG,"font":MONO,"size":25,"bold":True}),(" "+unit,{"color":MUTED,"size":12})],first=True)
    tf=box(s,x+0.28,4.75,3.35,1.4); para(tf,[(desc,{"color":MUTED,"size":12.5})],line=1.3,first=True)

# ===================================================== 4 · products / BOM
s=slide(); eyebrow(s,0.7,0.6,"03","PRODUCTS"); title(s,"Every part, to spec")
kicker(s,"A low-cost prototype proves the physics; a kiosk-grade production node ships to the transformer.")
def bom(s,x,y,w,head,col,rows,total):
    rrect(s,x,y,w,3.5,fill=PANEL,line=LINE,radius=0.05)
    tf=box(s,x+0.28,y+0.22,w-0.5,0.3); para(tf,[(head,{"color":col,"font":MONO,"size":11,"spacing":1.0})],first=True)
    tf=box(s,x+0.28,y+0.62,w-0.56,2.4)
    for j,(n,c) in enumerate(rows):
        para(tf,[(n,{"color":CY,"size":12.5}),("   "*1,{}),],first=(j==0),space_after=6)
        # price aligned right via separate box
    yy=y+0.66
    for (n,c) in rows:
        tfp=box(s,x+w-1.7,yy,1.4,0.3,anchor=MSO_ANCHOR.TOP); para(tfp,[(c,{"color":MUTED,"font":MONO,"size":11.5})],align=PP_ALIGN.RIGHT,first=True)
        yy+=0.365
    tf=box(s,x+0.28,y+3.02,w-0.56,0.4)
    para(tf,[(total[0],{"color":FG,"size":13,"bold":True})],first=True)
    tfp=box(s,x+w-2.2,y+3.0,1.9,0.4); para(tfp,[(total[1],{"color":col,"font":MONO,"size":16,"bold":True})],align=PP_ALIGN.RIGHT,first=True)
proto=[("ESP32 DevKit","EGP 250"),("SCT-013-000 CT clamp","EGP 350"),("Burden + bias circuit","EGP 60"),("ZMPT101B voltage sensor","EGP 150"),("Breadboard, PSU, jumpers","EGP 190")]
prod=[("Rogowski coil ×3","EGP 1,800"),("Metering AFE + MCU","EGP 600"),("LoRaWAN module (SX1262)","EGP 400"),("IP65 enclosure + power","EGP 1,200"),("Assembly, test & calibration","EGP 500")]
bom(s,0.7,2.7,5.9,"PROTOTYPE · PROOF OF CONCEPT",SIGNAL,proto,("Total prototype","≈ EGP 1,000"))
bom(s,6.75,2.7,5.9,"PRODUCTION NODE (hardware)",OK,prod,("+ install EGP 1,500  ·  all-in","≈ EGP 6,000"))
tf=box(s,0.7,6.5,11.9,0.4); para(tf,[("One instrumented transformer ≈ EGP 6,000 (~$125) — and it watches a whole feeder of customers.",{"color":DIM,"size":12})],first=True)

# ===================================================== 5 · unit economics
s=slide(); eyebrow(s,0.7,0.6,"04","UNIT ECONOMICS"); title(s,"One node, five-year life")
rows=[("All-in cost to deploy (CapEx)","EGP 6,000"),("Hardware gross margin (one-time)","EGP 1,500"),
      ("SaaS revenue / year","EGP 3,600"),("Recovery-share / year (blended)","EGP 1,900"),
      ("Recurring gross profit / year","EGP 3,850"),("Allocated CAC (acq. + install)","EGP 2,500")]
yy=2.7
for k,v in rows:
    tf=box(s,0.7,yy,4.6,0.35,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(k,{"color":MUTED,"size":13.5})],first=True)
    tfp=box(s,5.3,yy,1.6,0.35,anchor=MSO_ANCHOR.MIDDLE); para(tfp,[(v,{"color":FG,"font":MONO,"size":13.5})],align=PP_ALIGN.RIGHT,first=True)
    rect(s,0.7,yy+0.42,6.2,0.014,LINE); yy+=0.5
tf=box(s,0.7,yy+0.05,4.6,0.4); para(tf,[("5-yr gross profit / node (LTV)",{"color":FG,"size":14.5,"bold":True})],first=True)
tfp=box(s,4.9,yy+0.05,2.0,0.4); para(tfp,[("≈ EGP 20,750",{"color":OK,"font":MONO,"size":18,"bold":True})],align=PP_ALIGN.RIGHT,first=True)
tf=box(s,0.7,yy+0.6,4.6,0.4); para(tf,[("LTV / CAC",{"color":FG,"size":14.5,"bold":True})],first=True)
tfp=box(s,4.9,yy+0.6,2.0,0.4); para(tfp,[("≈ 8×",{"color":OK,"font":MONO,"size":18,"bold":True})],align=PP_ALIGN.RIGHT,first=True)
# customer ROI callout
rrect(s,7.5,2.7,5.1,3.7,fill=PANEL,line=LINE,radius=0.05); strip(s,7.5,2.7,3.7,OK)
tf=box(s,7.8,2.95,4.6,0.4); para(tf,[("THE CUSTOMER'S MATH",{"color":OK,"font":MONO,"size":11,"spacing":1.0})],first=True)
tf=box(s,7.8,3.35,4.6,0.9); para(tf,[("EGP 3,500",{"color":OK,"font":MONO,"size":40,"bold":True}),("  / mo",{"color":MUTED,"size":14})],first=True)
tf=box(s,7.8,4.4,4.5,1.8)
para(tf,[("recovered when a node catches one continuous ~3 kW illegal draw (≈2,160 kWh/mo).",{"color":CY,"size":13})],line=1.3,first=True)
para(tf,[("Against an EGP 6,000 node — ",{"color":CY,"size":13}),("payback in under two months.",{"color":OK,"size":13,"bold":True})],space_before=8,line=1.3)

# ===================================================== 6 · financials
s=slide(); eyebrow(s,0.7,0.6,"05","FINANCIALS"); title(s,"Five-year model — EGP millions")
# table (left)
cols=["","Y1","Y2","Y3","Y4","Y5"]
data=[("Total revenue",["4.4","33.5","142.8","324.4","604.8"],True),
      ("Gross profit",["1.9","15.0","67.0","167.1","327.8"],False),
      ("Operating expenses",["12.0","28.0","50.0","95.0","160.0"],False),
      ("EBITDA",["(10.1)","(13.0)","17.0","72.1","167.8"],True),
      ("Transformers (k)",["0.5","4","18","45","90"],False)]
tf=box(s,0.7,2.55,6.2,0.3)
p=tf.paragraphs[0]
for c in cols:
    r=p.add_run(); r.text=(c+"        ")[:9] if c=="" else c.rjust(8); r.font.name=MONO; r.font.size=Pt(11); r.font.color.rgb=SIGNAL; r.font.bold=True
yy=2.95
for name,vals,strong in data:
    tf=box(s,0.7,yy,2.4,0.32,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(name,{"color":FG if strong else MUTED,"size":12.5,"bold":strong})],first=True)
    for i,v in enumerate(vals):
        tfp=box(s,2.9+i*0.82,yy,0.78,0.32,anchor=MSO_ANCHOR.MIDDLE)
        col=THEFT if v.startswith("(") else (OK if (name=="EBITDA") else CY)
        para(tfp,[(v,{"color":col,"font":MONO,"size":12,"bold":strong})],align=PP_ALIGN.RIGHT,first=True)
    rect(s,0.7,yy+0.4,8.5,0.012,LINE); yy+=0.52
# stacked bar chart (right)
cx0,cy0,cw,ch=9.5,3.0,3.0,3.1
hw=[3,21,84,162,270]; saas=[0.9,8.1,39.6,113.4,243]; rec=[0.5,4.4,19.2,49,91.8]
tot=[h+saas[i]+rec[i] for i,h in enumerate(hw)]; mx=max(tot)
rect(s,cx0,cy0+ch,cw,0.012,LINE2)
bw=cw/5*0.55
for i in range(5):
    cxc=cx0+(i+0.5)*cw/5; base=cy0+ch
    for val,col in [(hw[i],SIGNAL),(saas[i],OK),(rec[i],TECH)]:
        hh=val/mx*ch
        rect(s,cxc-bw/2,base-hh,bw,hh,col); base-=hh
    tfp=box(s,cxc-0.4,cy0+ch+0.05,0.8,0.25); para(tfp,[("Y"+str(i+1),{"color":MUTED,"font":MONO,"size":10})],align=PP_ALIGN.CENTER,first=True)
tf=box(s,cx0,cy0-0.35,cw,0.3); para(tf,[("Revenue by stream",{"color":MUTED,"size":11})],first=True)
tf=box(s,9.5,6.55,3.2,0.3)
para(tf,[("■ ",{"color":SIGNAL,"size":10}),("Hardware  ",{"color":MUTED,"size":9.5}),("■ ",{"color":OK,"size":10}),("SaaS  ",{"color":MUTED,"size":9.5}),("■ ",{"color":TECH,"size":10}),("Recovery",{"color":MUTED,"size":9.5})],first=True)
tf=box(s,0.7,6.55,8.5,0.3); para(tf,[("Breakeven in Year 3 as recurring revenue overtakes the cost base.",{"color":DIM,"size":12})],first=True)

# ===================================================== 7 · deployment (photo)
s=slide(); eyebrow(s,0.7,0.6,"06","DEPLOYMENT"); title(s,"Onto the real iron box")
kicker(s,"The field target is a sealed distribution kiosk (RMU) in Kafr El-Sheikh — retrofitted non-invasively on the LV side.")
photo(s,"kiosk.jpg",0.9,2.75,3.7)
tf=box(s,0.9,6.05,3.7,0.3); para(tf,[("Distribution kiosk (RMU) · Kafr El-Sheikh",{"color":DIM,"font":MONO,"size":10})],first=True)
dp=[("Wrap 3 Rogowski coils","around the LV busbars — one per phase, no cables cut"),
    ("IP65 edge node","custom PCB in the panel; antenna outside the metal shell"),
    ("MV chamber never touched","magnetic coupling — no galvanic contact, no outage"),
    ("Scales to 90,000 kiosks","~30 min, reversible, tamper-resistant per unit")]
yy=2.95
for t,d in dp:
    tf=box(s,5.2,yy,7.4,0.8)
    para(tf,[("✓  ",{"color":OK,"font":MONO,"size":14,"bold":True}),(t,{"color":FG,"size":16,"bold":True})],first=True)
    para(tf,[("      "+d,{"color":MUTED,"size":13})],space_before=1)
    yy+=0.87

# ===================================================== competition
s=slide(); eyebrow(s,0.7,0.6,"07","COMPETITIVE LANDSCAPE"); title(s,"Where we sit")
kicker(s,"Six ways utilities fight theft today — only one is low-CapEx, explainable, and works without a smart-meter rollout.")
cxs=[4.5,6.1,7.6,9.1,10.7]
heads=["Needs AMI","Sep. theft","Explains","CapEx","Legacy grid"]
tf=box(s,0.75,2.62,3.4,0.3); para(tf,[("APPROACH",{"color":SIGNAL,"font":MONO,"size":10,"spacing":.6})],first=True)
for i,h in enumerate(heads):
    tf=box(s,cxs[i]-0.8,2.55,1.6,0.4); para(tf,[(h,{"color":SIGNAL,"font":MONO,"size":9.5})],align=PP_ALIGN.CENTER,first=True)
crows=[("Manual field audits",["No","x","partial","labour","yes"]),
       ("Full AMI rollout",["is meter","analytics","partial","very high","new only"]),
       ("AMI / MDM analytics",["Yes","check","check","high","needs AMI"]),
       ("LV feeder monitoring",["No","x","x","high","yes"]),
       ("Fixed-threshold alarms",["No","x","x","low","yes"]),
       ("Grid-Pulse NTL",["No","check","check","low","yes"])]
def cell(v):
    return {"check":("✓",OK),"x":("✗",THEFT),"yes":("✓",OK),"No":("No",OK)}.get(v,(v,MUTED))
yy=3.12
for name,cells in crows:
    us=(name=="Grid-Pulse NTL")
    if us: rrect(s,0.6,yy-0.05,12.05,0.5,fill=RGBColor(0x0F,0x1B,0x1E),line=SIGDIM,radius=0.1)
    tf=box(s,0.78,yy,3.5,0.4,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(name,{"color":(OK if us else FG),"size":12.5,"bold":us})],first=True)
    for i,v in enumerate(cells):
        txt,col=cell(v)
        tf=box(s,cxs[i]-0.8,yy,1.6,0.4,anchor=MSO_ANCHOR.MIDDLE)
        para(tf,[(txt,{"color":col,"font":MONO,"size":11.5,"bold":txt in("✓","✗")})],align=PP_ALIGN.CENTER,first=True)
    yy+=0.56
tf=box(s,0.7,yy+0.18,11.9,0.5)
para(tf,[("The whitespace:  ",{"color":FG,"size":13.5,"bold":True}),
         ("low-CapEx + explainable + no AMI required — the corner every incumbent misses.",{"color":MUTED,"size":13.5})],first=True)

# ===================================================== 8 · moat & risks
s=slide(); eyebrow(s,0.7,0.6,"08","MOAT & RISK"); title(s,"Why we win — and what could go wrong")
tf=box(s,0.7,2.5,6,0.3); para(tf,[("THE MOAT vs ALTERNATIVES",{"color":SIGNAL,"font":MONO,"size":10.5,"spacing":1.0})],first=True)
moat=[("Manual audits","random, slow"),("Smart-meter rollout","years · huge CapEx"),
      ("Fixed-threshold analytics","floods on hot days"),("Grid-Pulse NTL","localized, explained, ROI < 2 mo")]
yy=2.95
for i,(a,b) in enumerate(moat):
    us=(i==3); col=OK if us else CY
    tf=box(s,0.7,yy,2.7,0.34,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(a,{"color":(OK if us else FG),"size":13,"bold":us})],first=True)
    tf=box(s,3.5,yy,3.3,0.34,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(b,{"color":col,"size":12.5})],first=True)
    rect(s,0.7,yy+0.42,6.1,0.012,LINE); yy+=0.52
tf=box(s,7.2,2.5,5.4,0.3); para(tf,[("RISKS → MITIGATION",{"color":TECH,"font":MONO,"size":10.5,"spacing":1.0})],first=True)
risks=[("Slow B2G procurement","paid pilot + audited recovery → tender"),
       ("Hardware CapEx at scale","asset-backed financing vs signed contracts"),
       ("Recovery-share decays","shift to recurring SaaS + new regions"),
       ("Model false positives","explainable verdicts + human-in-loop")]
yy=2.95
for t,m in risks:
    tf=box(s,7.2,yy,5.4,0.75)
    para(tf,[("▲ ",{"color":TECH,"size":11}),(t,{"color":FG,"size":13,"bold":True})],first=True)
    para(tf,[("   ",{}),(m,{"color":MUTED,"size":11.5})],space_before=1)
    yy+=0.87

# ===================================================== 9 · the ask
s=slide(); eyebrow(s,0.7,0.6,"09","THE ASK"); title(s,"EGP 30M seed · use of funds")
tf=box(s,0.7,2.7,4.2,1.4)
para(tf,[("EGP 30M",{"color":SIGNAL,"font":MONO,"size":46,"bold":True})],first=True)
para(tf,[("≈ $625K — carries the company through the district pilot and two years to the Year-3 EBITDA turn.",{"color":MUTED,"size":13.5})],space_before=8,line=1.35)
uof=[("Field pilot — 1,000 nodes + gateways",40,"EGP 12.0M"),("Product & ML engineering",25,"EGP 7.5M"),
     ("Hardware R&D + certification",15,"EGP 4.5M"),("Team & G&A",12,"EGP 3.6M"),("Go-to-market / BD",8,"EGP 2.4M")]
yy=2.75
for n,p,v in uof:
    tf=box(s,5.3,yy,3.7,0.3,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(n,{"color":CY,"size":12.5})],first=True)
    rrect(s,9.1,yy+0.03,2.4,0.16,fill=RGBColor(0x0D,0x15,0x19),line=LINE,line_w=0.75,radius=0.5)
    rect(s,9.12,yy+0.05,2.36*p/100,0.12,SIGNAL)
    tf=box(s,11.6,yy,1.1,0.3,anchor=MSO_ANCHOR.MIDDLE); para(tf,[(v,{"color":MUTED,"font":MONO,"size":11})],align=PP_ALIGN.RIGHT,first=True)
    yy+=0.62
tf=box(s,0.7,6.5,11.9,0.5)
para(tf,[("Hardware at scale (Years 3–5) is funded by asset-backed financing against signed EDC contracts — keeping the equity raise small vs the EGP 600M+ revenue it unlocks.",{"color":DIM,"size":12})],line=1.3,first=True)

out="Grid-Pulse-NTL-Business.pptx"; prs.save(out)
print("saved",out,"-",len(prs.slides._sldIdLst),"slides")
